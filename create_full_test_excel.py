"""
Tạo file Excel test đầy đủ 8 loại câu hỏi, mỗi loại 2 câu.
Cột: No | Question Type | Question Text | Option A-D | Correct Answer | Explanation
"""
try:
    import openpyxl
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:
    import subprocess, sys
    subprocess.check_call([sys.executable, '-m', 'pip', 'install', 'openpyxl', '-q'])
    import openpyxl
    from openpyxl.worksheet.datavalidation import DataValidation

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = Workbook()
ws = wb.active
ws.title = "Câu hỏi"

# ── Headers ──────────────────────────────────────────────────────
HEADERS = [
    "Question No", "Question Type", "Question Text",
    "Option A", "Option B", "Option C", "Option D",
    "Correct Answer", "Explanation",
]
WIDTHS = [13, 30, 52, 24, 24, 24, 24, 22, 36]

TYPE_OPTIONS = [
    "Trắc nghiệm (1 đáp án)",
    "Trắc nghiệm (nhiều đáp án)",
    "Đúng / Sai",
    "Điền vào chỗ trống",
    "Viết ngắn",
    "Nối đôi",
    "Sắp xếp thứ tự",
    "Nghe hiểu",
]

# Màu riêng cho từng loại câu (cột B)
TYPE_COLORS = {
    "Trắc nghiệm (1 đáp án)":    "DBEAFE",  # xanh nhạt
    "Trắc nghiệm (nhiều đáp án)": "EDE9FE",  # tím nhạt
    "Đúng / Sai":                 "DCFCE7",  # xanh lá nhạt
    "Điền vào chỗ trống":         "FEF9C3",  # vàng nhạt
    "Viết ngắn":                  "FCE7F3",  # hồng nhạt
    "Nối đôi":                    "E0F2FE",  # xanh nhạt 2
    "Sắp xếp thứ tự":             "F3E8FF",  # tím nhạt 2
    "Nghe hiểu":                  "FFF7ED",  # cam nhạt
}

# ── Row 1: Header ────────────────────────────────────────────────
hdr_fill = PatternFill("solid", fgColor="1E3A8A")
for ci, h in enumerate(HEADERS, 1):
    c = ws.cell(row=1, column=ci, value=h)
    c.font      = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    c.fill      = hdr_fill
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.column_dimensions[get_column_letter(ci)].width = WIDTHS[ci - 1]
ws.row_dimensions[1].height = 28

# ── Dropdown data validation ──────────────────────────────────────
ws_lists = wb.create_sheet("_Lists")
ws_lists.sheet_state = "hidden"
for ri, opt in enumerate(TYPE_OPTIONS, 1):
    ws_lists.cell(row=ri, column=1, value=opt)

dv = DataValidation(
    type="list",
    formula1=f"_Lists!$A$1:$A${len(TYPE_OPTIONS)}",
    allow_blank=True,
    showDropDown=False,
    showErrorMessage=True,
    errorTitle="Loại không hợp lệ",
    error="Vui lòng chọn từ danh sách.",
)
dv.sqref = "B2:B200"
ws.add_data_validation(dv)

# ── 16 câu hỏi test ─────────────────────────────────────────────
# (No, Type, Text, A, B, C, D, Answer, Explanation)
ROWS = [
    # ── 1. Trắc nghiệm 1 đáp án ─────────────────────────────────
    ("1", "Trắc nghiệm (1 đáp án)",
     "Which sentence uses the Present Simple correctly?",
     "She go to school every day.",
     "She goes to school every day.",
     "She going to school every day.",
     "She gone to school every day.",
     "B",
     "3rd person singular: verb + 's' → goes."),

    ("2", "Trắc nghiệm (1 đáp án)",
     "What is the capital city of Australia?",
     "Sydney", "Melbourne", "Canberra", "Brisbane",
     "C",
     "Canberra is the capital, not Sydney."),

    # ── 2. Trắc nghiệm nhiều đáp án ─────────────────────────────
    ("3", "Trắc nghiệm (nhiều đáp án)",
     "Which of the following are vowels in English?",
     "A", "B", "E", "K",
     "A,C",
     "Vowels: A, E, I, O, U."),

    ("4", "Trắc nghiệm (nhiều đáp án)",
     "Which words are adjectives in the sentence: 'The tall, happy boy ran quickly'?",
     "tall", "boy", "happy", "quickly",
     "A,C",
     "Tall and happy describe the noun 'boy'."),

    # ── 3. Đúng / Sai ────────────────────────────────────────────
    ("5", "Đúng / Sai",
     "The Earth revolves around the Sun.",
     "", "", "", "",
     "Đúng",
     "The Earth orbits the Sun, not the other way around."),

    ("6", "Đúng / Sai",
     "Dolphins are a type of fish.",
     "", "", "", "",
     "Sai",
     "Dolphins are mammals, not fish."),

    # ── 4. Điền vào chỗ trống ────────────────────────────────────
    ("7", "Điền vào chỗ trống",
     "She ___ (study) English every evening. (Present Simple)",
     "", "", "", "",
     "studies",
     "3rd person singular: study → studies."),

    ("8", "Điền vào chỗ trống",
     "Water ___ (boil) at 100 degrees Celsius.",
     "", "", "", "",
     "boils",
     "Scientific fact uses Present Simple: boils."),

    # ── 5. Viết ngắn ─────────────────────────────────────────────
    ("9", "Viết ngắn",
     "Describe what you did last weekend in 2-3 sentences.",
     "", "", "", "",
     "",
     ""),

    ("10", "Viết ngắn",
     "What is the difference between 'much' and 'many'? Give one example each.",
     "", "", "", "",
     "",
     ""),

    # ── 6. Nối đôi ──────────────────────────────────────────────
    ("11", "Nối đôi",
     "Match each English word with its Vietnamese meaning:",
     "", "", "", "",
     "happy=vui vẻ, sad=buồn bã, angry=tức giận, surprised=ngạc nhiên",
     ""),

    ("12", "Nối đôi",
     "Match the verb tense with its usage:",
     "", "", "", "",
     "Present Simple=thói quen hàng ngày, Past Simple=hành động đã xảy ra, Future Simple=kế hoạch tương lai",
     ""),

    # ── 7. Sắp xếp thứ tự ───────────────────────────────────────
    ("13", "Sắp xếp thứ tự",
     "Rearrange the words to form a correct sentence:",
     "", "", "", "",
     "I, go, to, school, every, day",
     "Correct: I go to school every day."),

    ("14", "Sắp xếp thứ tự",
     "Put these historical events in chronological order:",
     "", "", "", "",
     "World War I begins, World War II begins, Moon landing, Berlin Wall falls",
     "1914, 1939, 1969, 1989."),

    # ── 8. Nghe hiểu ─────────────────────────────────────────────
    ("15", "Nghe hiểu",
     "Listen and choose: What does Tom order at the restaurant?",
     "Pizza", "Pasta", "Salad", "Soup",
     "B",
     "Tom says: 'I'll have the pasta, please.'"),

    ("16", "Nghe hiểu",
     "Listen and answer: How does the speaker get to work?",
     "By car", "By bus", "By bicycle", "On foot",
     "C",
     "Speaker says she cycles to work every morning."),
]

# ── Ghi dữ liệu ──────────────────────────────────────────────────
thin   = Side(style="thin", color="E2E8F0")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
alt    = [PatternFill("solid", fgColor="FFFFFF"), PatternFill("solid", fgColor="F8FAFC")]

for ri, row_data in enumerate(ROWS, 2):  # data bắt đầu row 2
    q_type = row_data[1]
    type_color = TYPE_COLORS.get(q_type, "FFFFFF")

    for ci, val in enumerate(row_data, 1):
        c = ws.cell(row=ri, column=ci, value=val)
        c.font      = Font(name="Calibri", size=10)
        c.alignment = Alignment(vertical="center", wrap_text=True)
        c.border    = border

        # Cột B (type) dùng màu riêng
        if ci == 2:
            c.fill = PatternFill("solid", fgColor=type_color)
            c.font = Font(name="Calibri", size=10, bold=True)
        else:
            c.fill = alt[ri % 2]

    ws.row_dimensions[ri].height = 26

ws.freeze_panes = "A2"

# ── Lưu ─────────────────────────────────────────────────────────
OUT = r"E:\IT_Project\01. ms Giang English\english-ms-giang\public\test-full-types.xlsx"
wb.save(OUT)
print(f"✅ Đã tạo: {OUT}")
print(f"\n{'No':<4} {'Loại':<32} Nội dung")
print("─" * 80)
for row in ROWS:
    print(f"{row[0]:<4} {row[1]:<32} {row[2][:42]}...")
