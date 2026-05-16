"""
Tạo file Excel mẫu cho CreateAssignmentDrawer.
Chạy: python create_excel_template.py
"""
import subprocess, sys

# Kiểm tra openpyxl
try:
    import openpyxl
except ImportError:
    subprocess.check_call([sys.executable, '-m', 'pip', 'install', 'openpyxl', '-q'])
    import openpyxl

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers
from openpyxl.utils import get_column_letter

wb = Workbook()
ws = wb.active
ws.title = "Câu hỏi"

# ── Màu sắc ──
HEADER_BG   = "1E40AF"   # Blue-800
HEADER_FG   = "FFFFFF"
REQ_BG      = "DBEAFE"   # Blue-100 (bắt buộc)
OPT_BG      = "F0FDF4"   # Green-50 (tuỳ chọn)
WARN_BG     = "FEF3C7"   # Amber-100 (cảnh báo)
ALT_ROW_BG  = "F8FAFC"

# ── Header ──
HEADERS = [
    ("Unit",            "A", 16, OPT_BG),
    ("Skill",           "B", 18, OPT_BG),
    ("Question No",     "C", 14, OPT_BG),
    ("Question Type",   "D", 22, REQ_BG),
    ("Question Text",   "E", 50, REQ_BG),
    ("Option A",        "F", 28, OPT_BG),
    ("Option B",        "G", 28, OPT_BG),
    ("Option C",        "H", 28, OPT_BG),
    ("Option D",        "I", 28, OPT_BG),
    ("Correct Answer",  "J", 20, WARN_BG),
    ("Audio File",      "K", 26, OPT_BG),
    ("Transcript",      "L", 36, OPT_BG),
    ("Explanation",     "M", 36, OPT_BG),
]

thin  = Side(style="thin",  color="CBD5E1")
thick = Side(style="medium", color="1E40AF")

for label, col, width, cell_bg in HEADERS:
    cell = ws[f"{col}1"]
    cell.value = label
    cell.font  = Font(name="Calibri", bold=True, color=HEADER_FG, size=11)
    cell.fill  = PatternFill("solid", fgColor=HEADER_BG)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = Border(
        left  = thick if col == "A" else thin,
        right = thick if col == "M" else thin,
        top   = thick, bottom=thick
    )
    ws.column_dimensions[col].width = width

ws.row_dimensions[1].height = 32

# ── Dữ liệu mẫu ──
SAMPLES = [
    # Unit, Skill, No, Type, Text, A, B, C, D, Answer, Audio, Transcript, Explanation
    ("Unit 2", "Grammar", "1", "multiple_choice",
     "What is the capital of England?",
     "Paris", "London", "Berlin", "Rome",
     "B", "", "", "London is the capital and largest city of England."),

    ("Unit 2", "Grammar", "2", "true_false",
     "The sun rises in the west.",
     "", "", "", "",
     "Sai", "", "", "The sun rises in the east, not the west."),

    ("Unit 2", "Vocabulary", "3", "fill_in",
     "She ___ (go) to school every day.",
     "", "", "", "",
     "goes", "", "", "Present simple: She + goes (3rd person singular)"),

    ("Unit 3", "Listening", "4", "listening",
     "Listen and answer: Where does the man want to go?",
     "The library", "The supermarket", "The hospital", "The school",
     "B", "audio_unit3_q4.mp3", "Man: Excuse me, could you tell me the way to the supermarket?", ""),

    ("Unit 3", "Writing", "5", "short_answer",
     "Write one sentence about your daily routine.",
     "", "", "", "",
     "", "", "", "Accept any grammatically correct sentence about daily routine."),
]

for i, row_data in enumerate(SAMPLES):
    row_num = i + 2
    bg = "FFFFFF" if i % 2 == 0 else ALT_ROW_BG
    for col_idx, val in enumerate(row_data):
        col_letter = get_column_letter(col_idx + 1)
        cell = ws[f"{col_letter}{row_num}"]
        cell.value = val
        cell.font  = Font(name="Calibri", size=10)
        cell.fill  = PatternFill("solid", fgColor=bg)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = Border(
            left=thin, right=thin, top=thin, bottom=thin
        )
    ws.row_dimensions[row_num].height = 28

# ── Freeze header row ──
ws.freeze_panes = "A2"

# ── Sheet hướng dẫn ──
guide = wb.create_sheet("Hướng dẫn")
guide.column_dimensions["A"].width = 24
guide.column_dimensions["B"].width = 60

guide_data = [
    ("📋 Hướng dẫn sử dụng", ""),
    ("", ""),
    ("Cột", "Mô tả"),
    ("Unit",           "Tên unit (tuỳ chọn, ví dụ: Unit 2)"),
    ("Skill",          "Kỹ năng (tuỳ chọn, ví dụ: Grammar, Listening)"),
    ("Question No",    "Số thứ tự câu (tuỳ chọn, hệ thống tự đánh số)"),
    ("Question Type",  "⚠️ BẮT BUỘC — xem các loại hợp lệ bên dưới"),
    ("Question Text",  "⚠️ BẮT BUỘC — nội dung câu hỏi"),
    ("Option A-D",     "Đáp án A, B, C, D (chỉ dùng cho multiple_choice)"),
    ("Correct Answer", "Đáp án đúng. Thiếu sẽ bị cảnh báo — vẫn import được"),
    ("Audio File",     "Tên file audio (chỉ cho Listening). Ví dụ: audio_unit2.mp3"),
    ("Transcript",     "Nội dung bài nghe (tuỳ chọn)"),
    ("Explanation",    "Giải thích đáp án (tuỳ chọn)"),
    ("", ""),
    ("✅ Loại câu hỏi hợp lệ", ""),
    ("multiple_choice","Trắc nghiệm 4 đáp án (A/B/C/D)"),
    ("true_false",     "Đúng / Sai"),
    ("fill_in",        "Điền vào chỗ trống"),
    ("short_answer",   "Trả lời ngắn"),
    ("matching",       "Nối từ"),
    ("listening",      "Nghe hiểu (cần Audio File)"),
    ("", ""),
    ("❌ Lỗi sẽ bị từ chối", ""),
    ("Thiếu Question Text",  "Dòng đó sẽ bị bỏ qua hoàn toàn"),
    ("Question Type sai",    "Dòng đó sẽ bị bỏ qua, kiểm tra chính tả"),
    ("⚠️ Cảnh báo (vẫn import)", ""),
    ("Thiếu Correct Answer", "Vẫn import, bổ sung ở bước Kiểm tra đáp án"),
    ("Listening thiếu Audio","Vẫn import, cần upload audio riêng sau"),
]

for r_idx, (a, b) in enumerate(guide_data, 1):
    ca = guide[f"A{r_idx}"]
    cb = guide[f"B{r_idx}"]
    ca.value = a
    cb.value = b
    if r_idx == 1:
        ca.font = Font(bold=True, size=14, color="1E40AF")
    elif a in ("Cột", "✅ Loại câu hỏi hợp lệ", "❌ Lỗi sẽ bị từ chối", "⚠️ Cảnh báo (vẫn import)"):
        ca.font = Font(bold=True, size=10, color="1E40AF")
        cb.font = Font(bold=True, size=10)
    ca.alignment = Alignment(vertical="top")
    cb.alignment = Alignment(vertical="top", wrap_text=True)
    guide.row_dimensions[r_idx].height = 20

OUT = r"E:\IT_Project\01. ms Giang English\english-ms-giang\public\mau-bai-tap.xlsx"
wb.save(OUT)
print(f"✅ Đã tạo file mẫu: {OUT}")
